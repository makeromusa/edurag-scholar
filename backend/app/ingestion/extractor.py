import os
import io
import logging
from typing import List, Dict, Any, Optional
import fitz  # PyMuPDF
from pptx import Presentation
import docx

logger = logging.getLogger(__name__)


class ExtractedPage:
    def __init__(self, page_number: int, content: str, title: Optional[str] = None, metadata: Optional[Dict[str, Any]] = None):
        self.page_number = page_number
        self.content = content.strip()
        self.title = title or f"Page {page_number}"
        self.metadata = metadata or {}

    def to_dict(self) -> Dict[str, Any]:
        return {
            "page_number": self.page_number,
            "title": self.title,
            "content": self.content,
            "metadata": self.metadata
        }


def extract_text_from_pdf(file_bytes: bytes, file_name: str) -> List[ExtractedPage]:
    """
    Extracts text per page/slide from a PDF file using PyMuPDF.
    Detects if a page is a scanned image with minimal text.
    """
    pages: List[ExtractedPage] = []
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    
    for i, page in enumerate(doc):
        page_num = i + 1
        text = page.get_text("text").replace("\x00", "").replace("\u0000", "").strip()
        
        # Extract potential slide title (often the first non-empty line or block)
        blocks = page.get_text("blocks")
        title = None
        if blocks:
            first_block_text = blocks[0][4].replace("\x00", "").replace("\u0000", "").strip()
            if len(first_block_text) < 100:
                title = first_block_text.replace("\n", " ")

        is_scanned = len(text) < 40 and len(page.get_images()) > 0
        
        # Automatic Vision OCR fallback for scanned slides / syllabus documents with minimal text
        if len(text) < 25 and len(page.get_images()) > 0:
            try:
                from app.core.config import settings
                from openai import OpenAI
                import base64
                if settings.AZURE_OPENAI_ENDPOINT and settings.AZURE_OPENAI_API_KEY:
                    pix = page.get_pixmap(dpi=120)
                    img_bytes = pix.tobytes("jpeg")
                    b64 = base64.b64encode(img_bytes).decode("utf-8")
                    client = OpenAI(base_url=settings.AZURE_OPENAI_ENDPOINT, api_key=settings.AZURE_OPENAI_API_KEY)
                    resp = client.chat.completions.create(
                        model=settings.AZURE_OPENAI_CHAT_DEPLOYMENT or "gpt-4.1-mini",
                        messages=[{
                            "role": "user",
                            "content": [
                                {"type": "text", "text": "Transcribe all text from this academic course document accurately."},
                                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}
                            ]
                        }],
                        max_tokens=800
                    )
                    ocr_text = resp.choices[0].message.content or ""
                    if len(ocr_text.strip()) > 10:
                        text = ocr_text.replace("\x00", "").replace("\u0000", "").strip()
                        title = title or text.split("\n")[0][:100]
            except Exception as e:
                logger.warning(f"Vision OCR fallback failed for page {page_num}: {e}")

        # Ensure page content is never empty
        if not text:
            text = f"Course material {file_name} — Page {page_num}"
        
        metadata = {
            "file_name": file_name,
            "is_scanned": is_scanned,
            "total_pages": len(doc)
        }
        
        pages.append(ExtractedPage(page_number=page_num, content=text, title=title, metadata=metadata))
        
    doc.close()
    return pages


def extract_text_from_pptx(file_bytes: bytes, file_name: str) -> List[ExtractedPage]:
    """
    Extracts text per slide from a PowerPoint presentation (.pptx).
    Extracts slide title, body shapes, tables, and speaker notes.
    """
    pages: List[ExtractedPage] = []
    prs = Presentation(io.BytesIO(file_bytes))
    
    for i, slide in enumerate(prs.slides):
        page_num = i + 1
        title = None
        slide_texts: List[str] = []
        
        # Check title
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()
            slide_texts.append(f"Title: {title}")
            
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
            elif shape.has_table:
                table_text = []
                for row in shape.table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_text.append(" | ".join(row_data))
                if table_text:
                    slide_texts.append("\n[Table Content]:\n" + "\n".join(table_text))
                    
        # Check notes slide
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_texts.append(f"[Presenter Notes]: {notes}")
                
        full_content = "\n".join(slide_texts)
        metadata = {
            "file_name": file_name,
            "total_pages": len(prs.slides),
            "source_type": "pptx"
        }
        
        pages.append(ExtractedPage(page_number=page_num, content=full_content, title=title, metadata=metadata))
        
    return pages


def extract_text_from_docx(file_bytes: bytes, file_name: str) -> List[ExtractedPage]:
    """
    Extracts text from Word documents (.docx) grouping by sections or headings.
    """
    doc = docx.Document(io.BytesIO(file_bytes))
    pages: List[ExtractedPage] = []
    current_title = "Document Overview"
    current_lines: List[str] = []
    page_counter = 1
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        if para.style.name.startswith("Heading"):
            if current_lines:
                pages.append(ExtractedPage(
                    page_number=page_counter,
                    content="\n".join(current_lines),
                    title=current_title,
                    metadata={"file_name": file_name, "source_type": "docx"}
                ))
                page_counter += 1
                current_lines = []
            current_title = text
            current_lines.append(f"## {text}")
        else:
            current_lines.append(text)
            
    if current_lines:
        pages.append(ExtractedPage(
            page_number=page_counter,
            content="\n".join(current_lines),
            title=current_title,
            metadata={"file_name": file_name, "source_type": "docx"}
        ))
        
    return pages


def extract_document_pages(file_bytes: bytes, file_name: str) -> List[ExtractedPage]:
    """
    Universal router for document extraction based on extension.
    """
    ext = os.path.splitext(file_name)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes, file_name)
    elif ext in [".pptx", ".ppt"]:
        return extract_text_from_pptx(file_bytes, file_name)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_bytes, file_name)
    else:
        # Fallback raw text
        try:
            raw_text = file_bytes.decode("utf-8")
            return [ExtractedPage(page_number=1, content=raw_text, title=file_name, metadata={"file_name": file_name})]
        except Exception:
            raise ValueError(f"Unsupported file format: {ext}")
